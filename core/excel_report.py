"""
สร้างไฟล์ Excel วิเคราะห์เชิงสถิติจากข้อมูลลูกค้าที่โทรแล้ว (ส่งวิชา Data Tool)
ทุกค่าในไฟล์เป็น "สูตร Excel จริง" อ้างอิงชีต Data แก้ข้อมูลแล้วผลคำนวณใหม่เอง

ชีตในไฟล์
  README        คำอธิบาย วิธีใช้ และสมมติฐาน
  Data          ข้อมูลลูกค้าที่โทรแล้ว + คอลัมน์ช่วยคำนวณ (สูตร)
  Descriptive   สถิติเชิงพรรณนา + ตารางแจกแจงความถี่อายุ + กราฟ
  Conversion    อัตราตอบรับแยกตามกลุ่ม (แบบ Pivot ด้วย COUNTIFS) + กราฟ
  T-Test        Welch's t-test เทียบกลุ่มตอบรับ vs ไม่ตอบรับ
  ANOVA         One-way ANOVA
  Chi-Square    Chi-square test of independence + Cramér's V
  Z-Test        เปรียบเทียบสัดส่วนสองกลุ่ม + ช่วงความเชื่อมั่นอัตราตอบรับ
  Correlation   เมทริกซ์สหสัมพันธ์ (CORREL)
  Solver        ปัญหาจัดสรรจำนวนสายโทร (Linear Programming) ตั้งค่า Solver ไว้ให้แล้ว
"""
from io import BytesIO

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from core import config as C

# ------------------------------------------------------------
# รูปแบบ
# ------------------------------------------------------------
FONT_NAME = "Tahoma"                       # รองรับภาษาไทย
F_BASE = Font(name=FONT_NAME, size=10)
F_BOLD = Font(name=FONT_NAME, size=10, bold=True)
F_TITLE = Font(name=FONT_NAME, size=14, bold=True, color="0F1E3D")
F_SUB = Font(name=FONT_NAME, size=10, italic=True, color="64748B")
F_HEAD = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
F_INPUT = Font(name=FONT_NAME, size=10, color="0000FF")       # ค่าที่ผู้ใช้แก้ได้
F_LINK = Font(name=FONT_NAME, size=10, color="008000")        # ดึงจากชีตอื่น
F_RESULT = Font(name=FONT_NAME, size=10, bold=True, color="0F1E3D")
FILL_HEAD = PatternFill("solid", fgColor="1D4ED8")
FILL_SECTION = PatternFill("solid", fgColor="E6EEFB")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_RESULT = PatternFill("solid", fgColor="E3F5F2")
THIN = Side(style="thin", color="D5DCE6")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PCT = "0.0%"
NUM2 = "#,##0.00"
NUM3 = "0.000"
NUM4 = "0.0000"
PVAL = "0.0000E+00"
INT = "#,##0"

# ห้ามใช้ "<30": COUNTIF จะตีความ "<" เป็นเครื่องหมายเปรียบเทียบ ไม่ใช่ชื่อกลุ่ม
AGE_GROUPS = ["17-29", "30-39", "40-49", "50-59", "60+"]
MONTHS = C.ALLOWED["month"]


def _style_range(ws, ref, font=None, fill=None, fmt=None, border=True, align=None):
    for row in ws[ref]:
        for c in row:
            if font:
                c.font = font
            if fill:
                c.fill = fill
            if fmt:
                c.number_format = fmt
            if border:
                c.border = BORDER
            if align:
                c.alignment = align


def _title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_SUB


def _section(ws, row, text, width=8):
    ws.cell(row=row, column=1, value=text).font = F_BOLD
    for col in range(1, width + 1):
        ws.cell(row=row, column=col).fill = FILL_SECTION


def _header(ws, row, labels, col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font = F_HEAD
        c.fill = FILL_HEAD
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _cell(ws, row, col, value, fmt=None, font=None, fill=None, border=True):
    c = ws.cell(row=row, column=col, value=value)
    c.font = font or F_BASE
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if border:
        c.border = BORDER
    return c


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ============================================================
# ชีต Data
# ============================================================
class DataRefs:
    """เก็บตำแหน่งคอลัมน์ในชีต Data เพื่อสร้างช่วงอ้างอิงแบบ absolute"""

    def __init__(self, columns, n_rows):
        self.col = {name: get_column_letter(i + 1) for i, name in enumerate(columns)}
        self.last = n_rows + 1

    def rng(self, name):
        letter = self.col[name]
        return f"Data!${letter}$2:${letter}${self.last}"


def _sheet_data(wb, df):
    ws = wb.create_sheet("Data")
    base_cols = [C.ID_COL] + C.DATA_COLUMNS
    helper_cols = ["target", "age_group", "age_yes", "age_no", "campaign_yes", "campaign_no"]
    columns = base_cols + helper_cols
    refs = DataRefs(columns, len(df))
    L = refs.col

    _header(ws, 1, columns)
    for i, rec in enumerate(df[base_cols].itertuples(index=False), start=2):
        values = [None if (isinstance(v, float) and pd.isna(v)) else
                  (v.item() if hasattr(v, "item") else v) for v in rec]
        ws.append(values)
        y, age, camp = f"{L['y']}{i}", f"{L['age']}{i}", f"{L['campaign']}{i}"
        ws[f"{L['target']}{i}"] = f'=IF({y}="yes",1,0)'
        ws[f"{L['age_group']}{i}"] = (f'=IF({age}<30,"17-29",IF({age}<40,"30-39",'
                                      f'IF({age}<50,"40-49",IF({age}<60,"50-59","60+"))))')
        ws[f"{L['age_yes']}{i}"] = f'=IF({y}="yes",{age},"")'
        ws[f"{L['age_no']}{i}"] = f'=IF({y}="no",{age},"")'
        ws[f"{L['campaign_yes']}{i}"] = f'=IF({y}="yes",{camp},"")'
        ws[f"{L['campaign_no']}{i}"] = f'=IF({y}="no",{camp},"")'

    ws.freeze_panes = "B2"
    for i, name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(18, len(name) + 3))
    first_helper = len(base_cols) + 1
    ws.cell(row=1, column=first_helper).comment = Comment(
        "คอลัมน์สีเขียวทางขวาเป็นสูตรช่วยคำนวณ: target (yes=1), กลุ่มอายุ, "
        "และแยกค่าอายุ/จำนวนครั้งที่โทรตามผลการโทร สำหรับ T.TEST", "Bank Call Tracker")
    for col in range(first_helper, len(columns) + 1):
        ws.cell(row=1, column=col).fill = PatternFill("solid", fgColor="0E9F8E")
    return refs


# ============================================================
# ชีต Descriptive
# ============================================================
NUMERIC_VARS = ["age", "duration", "campaign", "previous", "emp_var_rate",
                "cons_price_idx", "cons_conf_idx", "euribor3m", "nr_employed"]


def _sheet_descriptive(wb, refs):
    ws = wb.create_sheet("Descriptive")
    _title(ws, "สถิติเชิงพรรณนา (Descriptive Statistics)",
           "คำนวณจากลูกค้าที่โทรแล้วทั้งหมดในชีต Data | ทุกช่องเป็นสูตร")

    stats = [
        ("จำนวน (n)", "=COUNT({r})", INT),
        ("ค่าเฉลี่ย (Mean)", "=AVERAGE({r})", NUM2),
        ("มัธยฐาน (Median)", "=MEDIAN({r})", NUM2),
        ("ฐานนิยม (Mode)", '=IFERROR(_xlfn.MODE.SNGL({r}),"ไม่มีค่าซ้ำ")', NUM2),
        ("ส่วนเบี่ยงเบนมาตรฐาน (SD)", "=_xlfn.STDEV.S({r})", NUM2),
        ("ความแปรปรวน (Variance)", "=_xlfn.VAR.S({r})", NUM2),
        ("ค่าต่ำสุด (Min)", "=MIN({r})", NUM2),
        ("ควอร์ไทล์ที่ 1 (Q1)", "=_xlfn.QUARTILE.INC({r},1)", NUM2),
        ("ควอร์ไทล์ที่ 3 (Q3)", "=_xlfn.QUARTILE.INC({r},3)", NUM2),
        ("ค่าสูงสุด (Max)", "=MAX({r})", NUM2),
        ("พิสัย (Range)", "=MAX({r})-MIN({r})", NUM2),
        ("ความเบ้ (Skewness)", "=SKEW({r})", NUM3),
        ("ความโด่ง (Kurtosis)", "=KURT({r})", NUM3),
        ("ค่าเฉลี่ย กลุ่มตอบรับ", '=AVERAGEIFS({r},{y},"yes")', NUM2),
        ("ค่าเฉลี่ย กลุ่มไม่ตอบรับ", '=AVERAGEIFS({r},{y},"no")', NUM2),
    ]
    _header(ws, 4, ["สถิติ"] + [C.LABELS.get(v, v) for v in NUMERIC_VARS])
    for i, (label, tpl, fmt) in enumerate(stats, start=5):
        _cell(ws, i, 1, label, font=F_BOLD)
        for j, var in enumerate(NUMERIC_VARS, start=2):
            _cell(ws, i, j, tpl.format(r=refs.rng(var), y=refs.rng("y")), fmt=fmt)
    ci_row = 5 + len(stats)
    _cell(ws, ci_row, 1, "ช่วงความเชื่อมั่น 95% ของค่าเฉลี่ย (±)", font=F_BOLD)
    for j, var in enumerate(NUMERIC_VARS, start=2):
        r = refs.rng(var)
        _cell(ws, ci_row, j, f"=_xlfn.CONFIDENCE.T(0.05,_xlfn.STDEV.S({r}),COUNT({r}))", fmt=NUM3)
    _cell(ws, ci_row + 1, 1,
          "หมายเหตุ: duration (ความยาวสาย) รู้หลังโทรเสร็จ ใช้วิเคราะห์เชิงพรรณนาได้ แต่ห้ามใช้ทำนาย",
          font=F_SUB, border=False)

    # --- ตารางแจกแจงความถี่อายุ ---
    top = ci_row + 3
    _section(ws, top, "ตารางแจกแจงความถี่ของอายุ (Frequency Distribution)", 6)
    _header(ws, top + 1, ["ช่วงอายุ", "ขอบล่าง", "ขอบบน", "ความถี่", "ความถี่สัมพัทธ์", "ความถี่สะสม"])
    bins = [(17, 24), (25, 34), (35, 44), (45, 54), (55, 64), (65, 98)]
    age = refs.rng("age")
    first = top + 2
    for k, (lo, hi) in enumerate(bins):
        r = first + k
        _cell(ws, r, 1, f"{lo}-{hi}")
        _cell(ws, r, 2, lo, font=F_INPUT)
        _cell(ws, r, 3, hi, font=F_INPUT)
        _cell(ws, r, 4, f'=COUNTIFS({age},">="&B{r},{age},"<="&C{r})', fmt=INT)
        _cell(ws, r, 5, f"=D{r}/SUM($D${first}:$D${first + len(bins) - 1})", fmt=PCT)
        _cell(ws, r, 6, f"=SUM($D${first}:D{r})", fmt=INT)
    last = first + len(bins) - 1

    chart = BarChart()
    chart.type = "col"
    chart.title = "Histogram: Age"
    chart.y_axis.title = "Count"
    chart.add_data(Reference(ws, min_col=4, min_row=top + 1, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    chart.legend = None
    chart.height, chart.width = 7, 14
    ws.add_chart(chart, f"H{top}")

    _widths(ws, [34] + [15] * len(NUMERIC_VARS))
    ws.freeze_panes = "B5"


# ============================================================
# ชีต Conversion (แบบ Pivot)
# ============================================================
def _sheet_conversion(wb, refs, df):
    ws = wb.create_sheet("Conversion")
    _title(ws, "อัตราตอบรับแยกตามกลุ่ม (สรุปแบบ Pivot ด้วย COUNTIFS)",
           "อัตราตอบรับ = ตอบรับ / ทั้งหมด | Lift = อัตราตอบรับของกลุ่ม ÷ อัตราตอบรับรวม")
    y = refs.rng("y")
    _cell(ws, 4, 1, "อัตราตอบรับรวม", font=F_BOLD)
    _cell(ws, 4, 2, f'=COUNTIF({y},"yes")/COUNTA({y})', fmt=PCT, font=F_RESULT, fill=FILL_RESULT)
    overall = "$B$4"

    dims = [("age_group", AGE_GROUPS), ("job", None), ("marital", None), ("education", None),
            ("contact", None), ("month", None), ("poutcome", None)]
    row = 6
    for dim, order in dims:
        if order is None:
            present = df[dim].dropna().astype(str).unique().tolist()
            order = [m for m in MONTHS if m in present] if dim == "month" else \
                df[dim].value_counts().index.astype(str).tolist()
        _section(ws, row, f"แยกตาม {C.LABELS.get(dim, dim)}", 5)
        _header(ws, row + 1, [C.LABELS.get(dim, dim), "จำนวน", "ตอบรับ", "อัตราตอบรับ", "Lift"])
        g = refs.rng(dim)
        first = row + 2
        for k, val in enumerate(order):
            r = first + k
            _cell(ws, r, 1, val)
            _cell(ws, r, 2, f'=COUNTIF({g},A{r})', fmt=INT)
            _cell(ws, r, 3, f'=COUNTIFS({g},A{r},{y},"yes")', fmt=INT)
            _cell(ws, r, 4, f"=IF(B{r}=0,0,C{r}/B{r})", fmt=PCT)
            _cell(ws, r, 5, f"=IF({overall}=0,0,D{r}/{overall})", fmt="0.00x")
        last = first + len(order) - 1
        if dim in ("age_group", "job", "month", "contact"):
            ch = BarChart()
            ch.type = "col"
            ch.title = f"Conversion rate by {dim}"
            ch.add_data(Reference(ws, min_col=4, min_row=row + 1, max_row=last), titles_from_data=True)
            ch.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
            ch.legend = None
            ch.y_axis.numFmt = "0%"
            ch.height, ch.width = 6.5, 13
            ws.add_chart(ch, f"G{row}")
        row = max(last + 3, row + 15 if dim in ("age_group", "job", "month", "contact") else last + 3)

    ws.conditional_formatting.add(f"D6:D{row}", ColorScaleRule(
        start_type="min", start_color="FCE8E3", mid_type="percentile", mid_value=50,
        mid_color="FFFFFF", end_type="max", end_color="C6EFCE"))
    _widths(ws, [24, 12, 12, 14, 10])


# ============================================================
# ชีต T-Test
# ============================================================
def _sheet_ttest(wb, refs):
    ws = wb.create_sheet("T-Test")
    _title(ws, "Independent Two-Sample t-test (Welch's t-test)",
           "เทียบค่าเฉลี่ยของกลุ่ม 'ตอบรับ' กับ 'ไม่ตอบรับ' (ไม่สมมติว่าความแปรปรวนเท่ากัน)")
    _cell(ws, 4, 1, "ระดับนัยสำคัญ (α)", font=F_BOLD)
    _cell(ws, 4, 2, 0.05, fmt="0.00", font=F_INPUT, fill=FILL_INPUT).comment = Comment(
        "ค่าที่ผู้ใช้แก้ได้ (สมมติฐานมาตรฐาน α = 0.05)", "Bank Call Tracker")
    alpha = "$B$4"
    y = refs.rng("y")

    tests = [("age", "age_yes", "age_no", "อายุของคนที่ตอบรับ ต่างจากคนที่ไม่ตอบรับหรือไม่"),
             ("campaign", "campaign_yes", "campaign_no",
              "จำนวนครั้งที่โทรในแคมเปญ ของคนที่ตอบรับ ต่างจากคนที่ไม่ตอบรับหรือไม่")]
    row = 6
    for var, col_yes, col_no, question in tests:
        x = refs.rng(var)
        _section(ws, row, f"ทดสอบ: {C.LABELS[var]}", 4)
        _cell(ws, row + 1, 1, f"คำถาม: {question}", border=False)
        _cell(ws, row + 2, 1, "H₀: μ(ตอบรับ) = μ(ไม่ตอบรับ)    H₁: μ(ตอบรับ) ≠ μ(ไม่ตอบรับ)", border=False)
        h = row + 3
        _header(ws, h, ["", "ตอบรับ (yes)", "ไม่ตอบรับ (no)"])
        lines = [
            ("จำนวน (n)", f'=COUNTIF({y},"yes")', f'=COUNTIF({y},"no")', INT),
            ("ค่าเฉลี่ย", f'=AVERAGEIFS({x},{y},"yes")', f'=AVERAGEIFS({x},{y},"no")', NUM3),
            ("ความแปรปรวน", f'=SUMPRODUCT(({y}="yes")*({x}-B{h + 2})^2)/(B{h + 1}-1)',
             f'=SUMPRODUCT(({y}="no")*({x}-C{h + 2})^2)/(C{h + 1}-1)', NUM3),
        ]
        for k, (lab, fy, fn, fmt) in enumerate(lines, start=1):
            _cell(ws, h + k, 1, lab, font=F_BOLD)
            _cell(ws, h + k, 2, fy, fmt=fmt)
            _cell(ws, h + k, 3, fn, fmt=fmt)
        n1, n2, m1, m2, v1, v2 = (f"B{h + 1}", f"C{h + 1}", f"B{h + 2}", f"C{h + 2}",
                                  f"B{h + 3}", f"C{h + 3}")
        r = h + 5
        res = [
            ("ผลต่างค่าเฉลี่ย", f"={m1}-{m2}", NUM3),
            ("t statistic", f"=({m1}-{m2})/SQRT({v1}/{n1}+{v2}/{n2})", NUM3),
            ("องศาอิสระ (Welch df)",
             f"=({v1}/{n1}+{v2}/{n2})^2/(({v1}/{n1})^2/({n1}-1)+({v2}/{n2})^2/({n2}-1))", NUM2),
            ("p-value (สองหาง) จาก T.DIST.2T", f"=_xlfn.T.DIST.2T(ABS(B{r + 1}),B{r + 2})", PVAL),
            ("p-value ตรวจซ้ำด้วย T.TEST", f"=_xlfn.T.TEST({refs.rng(col_yes)},{refs.rng(col_no)},2,3)", PVAL),
            ("ค่าวิกฤต t (สองหาง)", f"=_xlfn.T.INV.2T({alpha},B{r + 2})", NUM3),
        ]
        for k, (lab, f, fmt) in enumerate(res):
            _cell(ws, r + k, 1, lab, font=F_BOLD)
            _cell(ws, r + k, 2, f, fmt=fmt)
        concl = r + len(res)
        _cell(ws, concl, 1, "สรุป", font=F_BOLD, fill=FILL_RESULT)
        _cell(ws, concl, 2, f'=IF(B{r + 3}<{alpha},"ปฏิเสธ H0: ค่าเฉลี่ยสองกลุ่มต่างกันอย่างมีนัยสำคัญ",'
                            f'"ไม่ปฏิเสธ H0: ไม่พบความแตกต่างอย่างมีนัยสำคัญ")', font=F_RESULT, fill=FILL_RESULT)
        ws.merge_cells(start_row=concl, start_column=2, end_row=concl, end_column=4)
        row = concl + 3
    _widths(ws, [36, 18, 18, 18])


# ============================================================
# ชีต ANOVA
# ============================================================
def _anova_block(ws, refs, row, var, factor, groups, alpha):
    x, g = refs.rng(var), refs.rng(factor)
    _section(ws, row, f"One-way ANOVA: {C.LABELS[var]} แยกตาม {C.LABELS[factor]}", 7)
    _cell(ws, row + 1, 1, f"H₀: ค่าเฉลี่ย{C.LABELS[var]}ของทุกกลุ่มเท่ากัน    "
                          "H₁: มีอย่างน้อย 1 กลุ่มที่ต่าง", border=False)
    h = row + 2
    _header(ws, h, ["กลุ่ม", "n", "ค่าเฉลี่ย", "ความแปรปรวน", "SS ภายในกลุ่ม", "SS ระหว่างกลุ่ม"])
    first = h + 1
    last = first + len(groups) - 1
    grand = f"$C${last + 2}"
    for k, grp in enumerate(groups):
        r = first + k
        _cell(ws, r, 1, grp)
        _cell(ws, r, 2, f"=COUNTIF({g},A{r})", fmt=INT)
        _cell(ws, r, 3, f"=AVERAGEIFS({x},{g},A{r})", fmt=NUM3)
        _cell(ws, r, 4, f"=IF(B{r}>1,E{r}/(B{r}-1),0)", fmt=NUM3)
        _cell(ws, r, 5, f"=SUMPRODUCT(({g}=A{r})*({x}-C{r})^2)", fmt=NUM2)
        _cell(ws, r, 6, f"=B{r}*(C{r}-{grand})^2", fmt=NUM2)
    tot = last + 1
    _cell(ws, tot, 1, "รวม", font=F_BOLD)
    _cell(ws, tot, 2, f"=SUM(B{first}:B{last})", fmt=INT, font=F_BOLD)
    _cell(ws, tot + 1, 1, "ค่าเฉลี่ยรวม (Grand mean)", font=F_BOLD)
    _cell(ws, tot + 1, 3, f"=SUMPRODUCT(B{first}:B{last},C{first}:C{last})/B{tot}", fmt=NUM3)

    t = tot + 3
    _header(ws, t, ["แหล่งความแปรปรวน", "SS", "df", "MS", "F", "p-value", "F วิกฤต"])
    k = len(groups)
    _cell(ws, t + 1, 1, "ระหว่างกลุ่ม (Between)")
    _cell(ws, t + 1, 2, f"=SUM(F{first}:F{last})", fmt=NUM2)
    _cell(ws, t + 1, 3, f"={k}-1", fmt=INT)
    _cell(ws, t + 1, 4, f"=B{t + 1}/C{t + 1}", fmt=NUM2)
    _cell(ws, t + 1, 5, f"=D{t + 1}/D{t + 2}", fmt=NUM3)
    _cell(ws, t + 1, 6, f"=_xlfn.F.DIST.RT(E{t + 1},C{t + 1},C{t + 2})", fmt=PVAL)
    _cell(ws, t + 1, 7, f"=_xlfn.F.INV.RT({alpha},C{t + 1},C{t + 2})", fmt=NUM3)
    _cell(ws, t + 2, 1, "ภายในกลุ่ม (Within)")
    _cell(ws, t + 2, 2, f"=SUM(E{first}:E{last})", fmt=NUM2)
    _cell(ws, t + 2, 3, f"=B{tot}-{k}", fmt=INT)
    _cell(ws, t + 2, 4, f"=B{t + 2}/C{t + 2}", fmt=NUM2)
    _cell(ws, t + 3, 1, "รวม (Total)", font=F_BOLD)
    _cell(ws, t + 3, 2, f"=B{t + 1}+B{t + 2}", fmt=NUM2)
    _cell(ws, t + 3, 3, f"=B{tot}-1", fmt=INT)
    _cell(ws, t + 4, 1, "η² (Effect size = SSB/SST)", font=F_BOLD)
    _cell(ws, t + 4, 2, f"=B{t + 1}/B{t + 3}", fmt=NUM4)
    _cell(ws, t + 5, 1, "สรุป", font=F_BOLD, fill=FILL_RESULT)
    _cell(ws, t + 5, 2, f'=IF(F{t + 1}<{alpha},"ปฏิเสธ H0: ค่าเฉลี่ยอย่างน้อย 1 กลุ่มต่างกันอย่างมีนัยสำคัญ",'
                        f'"ไม่ปฏิเสธ H0: ไม่พบความแตกต่างอย่างมีนัยสำคัญ")', font=F_RESULT, fill=FILL_RESULT)
    ws.merge_cells(start_row=t + 5, start_column=2, end_row=t + 5, end_column=7)
    return t + 8


def _sheet_anova(wb, refs, df):
    ws = wb.create_sheet("ANOVA")
    _title(ws, "One-way ANOVA (การวิเคราะห์ความแปรปรวนทางเดียว)",
           "SS ระหว่างกลุ่ม = Σ nᵢ(x̄ᵢ − x̄)²   SS ภายในกลุ่ม = ΣΣ(x − x̄ᵢ)²   F = MS ระหว่าง ÷ MS ภายใน")
    _cell(ws, 4, 1, "ระดับนัยสำคัญ (α)", font=F_BOLD)
    _cell(ws, 4, 2, 0.05, fmt="0.00", font=F_INPUT, fill=FILL_INPUT)
    marital = [m for m in ["married", "single", "divorced"] if m in set(df["marital"])]
    jobs = df["job"].value_counts().index.astype(str).tolist()
    nxt = _anova_block(ws, refs, 6, "age", "marital", marital, "$B$4")
    _anova_block(ws, refs, nxt, "campaign", "job", jobs, "$B$4")
    _widths(ws, [30, 16, 10, 16, 16, 16, 12])


# ============================================================
# ชีต Chi-Square
# ============================================================
def _chi_block(ws, refs, row, factor, cats, alpha):
    g, y = refs.rng(factor), refs.rng("y")
    k = len(cats)
    _section(ws, row, f"Chi-square: {C.LABELS.get(factor, factor)} × ผลการโทร", 9)
    _cell(ws, row + 1, 1, f"H₀: {C.LABELS.get(factor, factor)} กับผลการโทรเป็นอิสระต่อกัน (ไม่เกี่ยวข้องกัน)",
          border=False)
    h = row + 2
    _header(ws, h, [C.LABELS.get(factor, factor), "สังเกต: yes", "สังเกต: no", "รวม",
                    "คาดหวัง: yes", "คาดหวัง: no", "(O−E)²/E yes", "(O−E)²/E no"])
    first, last = h + 1, h + k
    tot = last + 1
    for i, cat in enumerate(cats):
        r = first + i
        _cell(ws, r, 1, cat)
        _cell(ws, r, 2, f'=COUNTIFS({g},A{r},{y},"yes")', fmt=INT)
        _cell(ws, r, 3, f'=COUNTIFS({g},A{r},{y},"no")', fmt=INT)
        _cell(ws, r, 4, f"=B{r}+C{r}", fmt=INT)
        _cell(ws, r, 5, f"=D{r}*B${tot}/$D${tot}", fmt=NUM2)
        _cell(ws, r, 6, f"=D{r}*C${tot}/$D${tot}", fmt=NUM2)
        _cell(ws, r, 7, f"=IF(E{r}=0,0,(B{r}-E{r})^2/E{r})", fmt=NUM3)
        _cell(ws, r, 8, f"=IF(F{r}=0,0,(C{r}-F{r})^2/F{r})", fmt=NUM3)
    _cell(ws, tot, 1, "รวม", font=F_BOLD)
    for col in "BCD":
        _cell(ws, tot, " BCD".index(col) + 1, f"=SUM({col}{first}:{col}{last})", fmt=INT, font=F_BOLD)

    s = tot + 2
    res = [
        ("χ² (Chi-square)", f"=SUM(G{first}:H{last})", NUM3),
        ("องศาอิสระ df = (r−1)(c−1)", f"=({k}-1)*(2-1)", INT),
        ("p-value จาก CHISQ.DIST.RT", f"=_xlfn.CHISQ.DIST.RT(B{s},B{s + 1})", PVAL),
        ("p-value ตรวจซ้ำด้วย CHISQ.TEST", f"=_xlfn.CHISQ.TEST(B{first}:C{last},E{first}:F{last})", PVAL),
        ("ค่าวิกฤต χ²", f"=_xlfn.CHISQ.INV.RT({alpha},B{s + 1})", NUM3),
        ("Cramér's V (ขนาดความสัมพันธ์)", f"=SQRT(B{s}/(D{tot}*(MIN({k},2)-1)))", NUM4),
        ("ช่องที่ค่าคาดหวัง < 5", f'=COUNTIF(E{first}:F{last},"<5")', INT),
    ]
    for i, (lab, f, fmt) in enumerate(res):
        _cell(ws, s + i, 1, lab, font=F_BOLD)
        _cell(ws, s + i, 2, f, fmt=fmt)
    concl = s + len(res)
    _cell(ws, concl, 1, "สรุป", font=F_BOLD, fill=FILL_RESULT)
    _cell(ws, concl, 2, f'=IF(B{s + 2}<{alpha},"ปฏิเสธ H0: เกี่ยวข้องกับผลการโทรอย่างมีนัยสำคัญ",'
                        f'"ไม่ปฏิเสธ H0: ไม่พบความเกี่ยวข้องอย่างมีนัยสำคัญ")', font=F_RESULT, fill=FILL_RESULT)
    ws.merge_cells(start_row=concl, start_column=2, end_row=concl, end_column=8)
    return concl + 3


def _sheet_chi(wb, refs, df):
    ws = wb.create_sheet("Chi-Square")
    _title(ws, "Chi-square Test of Independence",
           "ค่าคาดหวัง E = (ผลรวมแถว × ผลรวมคอลัมน์) ÷ ผลรวมทั้งหมด   χ² = Σ (O − E)² / E")
    _cell(ws, 4, 1, "ระดับนัยสำคัญ (α)", font=F_BOLD)
    _cell(ws, 4, 2, 0.05, fmt="0.00", font=F_INPUT, fill=FILL_INPUT)
    row = 6
    for factor in ["job", "contact", "poutcome", "marital"]:
        cats = df[factor].value_counts().index.astype(str).tolist()
        row = _chi_block(ws, refs, row, factor, cats, "$B$4")
    ages = df["age"].astype(int)
    present = [g for g, (lo, hi) in zip(AGE_GROUPS, [(0, 29), (30, 39), (40, 49), (50, 59), (60, 200)])
               if ages.between(lo, hi).any()]                  # เฉพาะกลุ่มที่มีข้อมูล
    row = _chi_block(ws, refs, row, "age_group", present, "$B$4")
    _widths(ws, [32, 14, 14, 12, 14, 14, 14, 14])


# ============================================================
# ชีต Z-Test (สัดส่วน)
# ============================================================
def _sheet_ztest(wb, refs):
    ws = wb.create_sheet("Z-Test")
    _title(ws, "Two-proportion z-test และช่วงความเชื่อมั่นของอัตราตอบรับ",
           "ใช้เปรียบเทียบอัตราตอบรับ (สัดส่วน) ระหว่างสองกลุ่ม")
    y, contact = refs.rng("y"), refs.rng("contact")
    _cell(ws, 4, 1, "ระดับนัยสำคัญ (α)", font=F_BOLD)
    _cell(ws, 4, 2, 0.05, fmt="0.00", font=F_INPUT, fill=FILL_INPUT)

    _section(ws, 6, "ทดสอบ: อัตราตอบรับ มือถือ (cellular) vs โทรศัพท์บ้าน (telephone)", 4)
    _cell(ws, 7, 1, "H₀: p(cellular) = p(telephone)    H₁: p(cellular) ≠ p(telephone)", border=False)
    _header(ws, 8, ["", "cellular", "telephone"])
    lines = [("จำนวน (n)", 'COUNTIF({c},"{v}")', INT),
             ("ตอบรับ (x)", 'COUNTIFS({c},"{v}",{y},"yes")', INT),
             ("อัตราตอบรับ (p̂)", None, PCT)]
    for i, (lab, tpl, fmt) in enumerate(lines, start=9):
        _cell(ws, i, 1, lab, font=F_BOLD)
        for j, v in enumerate(["cellular", "telephone"], start=2):
            col = get_column_letter(j)
            f = f"={tpl.format(c=contact, v=v, y=y)}" if tpl else f"=IF({col}9=0,0,{col}10/{col}9)"
            _cell(ws, i, j, f, fmt=fmt)
    res = [
        ("สัดส่วนรวม (pooled p)", "=(B10+C10)/(B9+C9)", PCT),
        ("z statistic", "=(B11-C11)/SQRT(B13*(1-B13)*(1/B9+1/C9))", NUM3),
        ("p-value (สองหาง)", "=2*(1-_xlfn.NORM.S.DIST(ABS(B14),TRUE))", PVAL),
        ("ค่าวิกฤต z", "=_xlfn.NORM.S.INV(1-$B$4/2)", NUM3),
    ]
    for i, (lab, f, fmt) in enumerate(res, start=13):
        _cell(ws, i, 1, lab, font=F_BOLD)
        _cell(ws, i, 2, f, fmt=fmt)
    _cell(ws, 17, 1, "สรุป", font=F_BOLD, fill=FILL_RESULT)
    _cell(ws, 17, 2, '=IF(B15<$B$4,"ปฏิเสธ H0: อัตราตอบรับสองช่องทางต่างกันอย่างมีนัยสำคัญ",'
                     '"ไม่ปฏิเสธ H0: ไม่พบความแตกต่างอย่างมีนัยสำคัญ")', font=F_RESULT, fill=FILL_RESULT)
    ws.merge_cells("B17:D17")

    _section(ws, 20, "ช่วงความเชื่อมั่นของอัตราตอบรับรวม", 4)
    ci = [("จำนวนที่โทร (n)", f"=COUNTA({y})", INT),
          ("อัตราตอบรับ (p̂)", f'=COUNTIF({y},"yes")/B21', PCT),
          ("ระดับความเชื่อมั่น", "=1-$B$4", "0%"),
          ("ค่าคลาดเคลื่อน (Margin of error)", "=_xlfn.NORM.S.INV(1-$B$4/2)*SQRT(B22*(1-B22)/B21)", PCT),
          ("ขอบล่าง", "=B22-B24", PCT),
          ("ขอบบน", "=B22+B24", PCT)]
    for i, (lab, f, fmt) in enumerate(ci, start=21):
        _cell(ws, i, 1, lab, font=F_BOLD)
        _cell(ws, i, 2, f, fmt=fmt)
    _widths(ws, [36, 16, 16, 16])


# ============================================================
# ชีต Correlation
# ============================================================
CORR_VARS = ["age", "duration", "campaign", "previous", "emp_var_rate", "cons_price_idx",
             "cons_conf_idx", "euribor3m", "nr_employed", "target"]


def _sheet_corr(wb, refs):
    ws = wb.create_sheet("Correlation")
    _title(ws, "เมทริกซ์สหสัมพันธ์ (Pearson Correlation)",
           "CORREL ระหว่างตัวแปรตัวเลข | target = 1 ถ้าตอบรับ (ค่ากับ target คือ point-biserial correlation)")
    labels = [C.LABELS.get(v, "ผลการโทร (target)") if v != "target" else "ผลการโทร (target)"
              for v in CORR_VARS]
    _header(ws, 4, [""] + labels)
    for i, a in enumerate(CORR_VARS):
        r = 5 + i
        _cell(ws, r, 1, labels[i], font=F_BOLD)
        for j, b in enumerate(CORR_VARS):
            _cell(ws, r, 2 + j, f"=CORREL({refs.rng(a)},{refs.rng(b)})", fmt="0.00")
    end = get_column_letter(1 + len(CORR_VARS))
    ws.conditional_formatting.add(f"B5:{end}{4 + len(CORR_VARS)}", ColorScaleRule(
        start_type="num", start_value=-1, start_color="F8696B", mid_type="num", mid_value=0,
        mid_color="FFFFFF", end_type="num", end_value=1, end_color="5A8AC6"))
    note = 6 + len(CORR_VARS)
    _cell(ws, note, 1, "อ่านค่า: ใกล้ +1 = สัมพันธ์ทางบวกสูง, ใกล้ −1 = ทางลบสูง, ใกล้ 0 = แทบไม่สัมพันธ์ "
                       "| ตัวแปรเศรษฐกิจสัมพันธ์กันเองสูงมาก (multicollinearity)", font=F_SUB, border=False)
    _widths(ws, [30] + [13] * len(CORR_VARS))


# ============================================================
# ชีต Solver
# ============================================================
def _sheet_solver(wb, refs):
    ws = wb.create_sheet("Solver")
    _title(ws, "Solver: จัดสรรจำนวนสายโทรให้กำไรคาดหวังสูงสุด (Linear Programming)",
           "ตัดสินใจว่าจะโทรหาลูกค้าแต่ละกลุ่ม (กลุ่มอายุ × ช่องทาง) กี่สาย ภายใต้งบจำนวนสายที่มี")

    _section(ws, 4, "ค่าสมมติ (แก้ได้ ช่องสีเหลือง)", 3)
    inputs = [("กำไรต่อคนที่ตอบรับ", C.PROFIT_PER_YES, INT, "ค่าสมมติ เท่ากับที่ใช้ในโมเดล ML"),
              ("ต้นทุนต่อสาย", C.COST_PER_CALL, INT, "ค่าสมมติ รวมเวลาของเซลส์"),
              ("งบจำนวนสายทั้งหมด", 800, INT, "ค่าสมมติ: ทีมขายโทรได้กี่สายในแคมเปญนี้"),
              ("โทรขั้นต่ำต่อกลุ่ม (% ของกลุ่ม)", 0.05, PCT,
               "ค่าสมมติ: นโยบายให้ทุกกลุ่มได้รับการติดต่ออย่างน้อยกี่ %")]
    for i, (lab, val, fmt, note) in enumerate(inputs, start=5):
        _cell(ws, i, 1, lab, font=F_BOLD)
        c = _cell(ws, i, 2, val, fmt=fmt, font=F_INPUT, fill=FILL_INPUT)
        c.comment = Comment(note, "Bank Call Tracker")
        _cell(ws, i, 3, note, font=F_SUB, border=False)
    _cell(ws, 9, 1, "จุดคุ้มทุน (อัตราตอบรับขั้นต่ำที่โทรแล้วคุ้ม)", font=F_BOLD)
    _cell(ws, 9, 2, "=B6/B5", fmt=PCT)

    _section(ws, 11, "ตารางตัดสินใจ", 9)
    _header(ws, 12, ["กลุ่มอายุ", "ช่องทาง", "ลูกค้าในกลุ่ม", "อัตราตอบรับ", "กำไรคาดหวังต่อสาย",
                     "จำนวนสายที่โทร (ตัวแปรตัดสินใจ)", "ขั้นต่ำ", "ผู้ตอบรับคาดหวัง", "กำไรคาดหวัง"])
    age, contact, y = refs.rng("age_group"), refs.rng("contact"), refs.rng("y")
    segs = [(a, c) for a in AGE_GROUPS for c in ["cellular", "telephone"]]
    first = 13
    last = first + len(segs) - 1
    for i, (a, c) in enumerate(segs):
        r = first + i
        _cell(ws, r, 1, a)
        _cell(ws, r, 2, c)
        _cell(ws, r, 3, f"=COUNTIFS({age},A{r},{contact},B{r})", fmt=INT)
        _cell(ws, r, 4, f'=IF(C{r}=0,0,COUNTIFS({age},A{r},{contact},B{r},{y},"yes")/C{r})', fmt=PCT)
        _cell(ws, r, 5, f"=D{r}*$B$5-$B$6", fmt=NUM2)
        _cell(ws, r, 6, f"=ROUND($B$7/{len(segs)},0)", fmt="#,##0.0", font=F_INPUT, fill=FILL_INPUT)
        _cell(ws, r, 7, f"=ROUNDUP(C{r}*$B$8,0)", fmt=INT)
        _cell(ws, r, 8, f"=F{r}*D{r}", fmt="#,##0.0")
        _cell(ws, r, 9, f"=F{r}*E{r}", fmt=NUM2)
    tot = last + 1
    _cell(ws, tot, 1, "รวม", font=F_BOLD)
    _cell(ws, tot, 6, f"=SUM(F{first}:F{last})", fmt="#,##0.0", font=F_BOLD)
    _cell(ws, tot, 8, f"=SUM(H{first}:H{last})", fmt="#,##0.0", font=F_BOLD)
    _cell(ws, tot, 9, f"=SUM(I{first}:I{last})", fmt=NUM2, font=F_RESULT, fill=FILL_RESULT)
    _cell(ws, tot + 1, 1, "เป้าหมาย (Objective)", font=F_BOLD)
    _cell(ws, tot + 1, 9, f"=I{tot}", fmt=NUM2, font=F_RESULT, fill=FILL_RESULT)
    _cell(ws, tot + 2, 1, "ตรวจข้อจำกัด: ใช้สายไม่เกินงบ", font=F_BOLD)
    _cell(ws, tot + 2, 6, f'=IF(F{tot}<=$B$7+0.0001,"ผ่าน","เกินงบ")')

    g = tot + 4
    _section(ws, g, "วิธีใช้ Solver (Data → Solver; ถ้าไม่มีเมนู ให้เปิด File → Options → Add-ins → Solver Add-in)", 9)
    steps = [
        f"1) Set Objective: $I${tot + 1}  → เลือก Max",
        f"2) By Changing Variable Cells: $F${first}:$F${last}",
        f"3) Add Constraint: $F${tot} <= $B$7   (ใช้สายไม่เกินงบ)",
        f"4) Add Constraint: $F${first}:$F${last} <= $C${first}:$C${last}   (โทรไม่เกินจำนวนลูกค้าในกลุ่ม)",
        f"5) Add Constraint: $F${first}:$F${last} >= $G${first}:$G${last}   (โทรขั้นต่ำต่อกลุ่ม)",
        "6) ติ๊ก Make Unconstrained Variables Non-Negative, Solving Method = Simplex LP → Solve",
        "ไฟล์นี้บันทึกการตั้งค่า Solver ไว้แล้ว เปิด Solver จะเห็นค่าข้างบนครบ กด Solve ได้ทันที",
        "ผลที่คาดว่าจะได้: Solver จะทุ่มสายไปที่กลุ่มที่ 'กำไรคาดหวังต่อสาย' สูงสุดก่อน "
        "และให้กลุ่มที่ขาดทุนต่อสายแค่ขั้นต่ำ",
    ]
    for i, s in enumerate(steps, start=g + 1):
        _cell(ws, i, 1, s, border=False)

    chart = BarChart()
    chart.type = "col"
    chart.title = "Calls allocated per segment"
    chart.add_data(Reference(ws, min_col=6, min_row=12, max_row=last), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, max_col=2, min_row=first, max_row=last))
    chart.legend = None
    chart.height, chart.width = 7, 16
    ws.add_chart(chart, f"K{first}")
    _widths(ws, [40, 14, 14, 13, 16, 18, 10, 16, 16])

    # --- บันทึกการตั้งค่า Solver ไว้ในไฟล์ (Excel อ่านจาก defined names ระดับชีต) ---
    sheet = "Solver"
    names = {
        "solver_opt": f"{sheet}!$I${tot + 1}", "solver_typ": "1", "solver_val": "0",
        "solver_adj": f"{sheet}!$F${first}:$F${last}",
        "solver_lhs1": f"{sheet}!$F${tot}", "solver_rel1": "1", "solver_rhs1": f"{sheet}!$B$7",
        "solver_lhs2": f"{sheet}!$F${first}:$F${last}", "solver_rel2": "1",
        "solver_rhs2": f"{sheet}!$C${first}:$C${last}",
        "solver_lhs3": f"{sheet}!$F${first}:$F${last}", "solver_rel3": "3",
        "solver_rhs3": f"{sheet}!$G${first}:$G${last}",
        "solver_num": "3", "solver_eng": "2", "solver_lin": "1", "solver_neg": "1", "solver_ver": "3",
    }
    for key, val in names.items():
        dn = DefinedName(key, attr_text=val, hidden=True)
        wb[sheet].defined_names[key] = dn


# ============================================================
# ชีต README
# ============================================================
def _sheet_readme(wb, n_rows, group_name, generated_at):
    ws = wb.active
    ws.title = "README"
    _title(ws, f"การวิเคราะห์เชิงสถิติ: แคมเปญเสนอขายเงินฝากประจำ ({group_name})",
           f"สร้างจากเว็บ Bank Call Tracker เมื่อ {generated_at} | ลูกค้าที่โทรแล้ว {n_rows:,} คน")
    rows = [
        ("ชีต", "เนื้อหา", "ฟังก์ชัน Excel หลัก"),
        ("Data", "ข้อมูลลูกค้าที่โทรแล้ว (ตอบรับ/ไม่ตอบรับ) + คอลัมน์สูตรช่วยคำนวณ", "IF"),
        ("Descriptive", "สถิติเชิงพรรณนา 9 ตัวแปร, ช่วงความเชื่อมั่น, ตารางแจกแจงความถี่ + Histogram",
         "AVERAGE, MEDIAN, MODE.SNGL, STDEV.S, VAR.S, QUARTILE.INC, SKEW, KURT, CONFIDENCE.T, COUNTIFS"),
        ("Conversion", "อัตราตอบรับและ Lift แยก 7 มิติ (สรุปแบบ Pivot) + กราฟ", "COUNTIF, COUNTIFS"),
        ("T-Test", "Welch's t-test: อายุ และจำนวนครั้งที่โทร ของกลุ่มตอบรับ vs ไม่ตอบรับ",
         "AVERAGEIFS, SUMPRODUCT, T.DIST.2T, T.TEST, T.INV.2T"),
        ("ANOVA", "One-way ANOVA: อายุตามสถานภาพ, จำนวนครั้งที่โทรตามอาชีพ + η²",
         "SUMPRODUCT, F.DIST.RT, F.INV.RT"),
        ("Chi-Square", "Chi-square test of independence 5 ตัวแปร × ผลการโทร + Cramér's V",
         "COUNTIFS, CHISQ.DIST.RT, CHISQ.TEST, CHISQ.INV.RT"),
        ("Z-Test", "Two-proportion z-test (มือถือ vs โทรศัพท์บ้าน) + ช่วงความเชื่อมั่นอัตราตอบรับ",
         "NORM.S.DIST, NORM.S.INV"),
        ("Correlation", "เมทริกซ์สหสัมพันธ์ 10 ตัวแปร + แถบสี", "CORREL"),
        ("Solver", "Linear Programming จัดสรรจำนวนสายโทรให้กำไรคาดหวังสูงสุด (ตั้งค่า Solver ไว้ให้แล้ว)",
         "SUMPRODUCT, Solver (Simplex LP)"),
    ]
    for i, row in enumerate(rows, start=4):
        for j, v in enumerate(row, start=1):
            c = _cell(ws, i, j, v, font=F_HEAD if i == 4 else F_BASE, fill=FILL_HEAD if i == 4 else None)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    r = 4 + len(rows) + 1
    notes = [
        "สีของช่อง: ตัวอักษรสีน้ำเงิน + พื้นเหลือง = ค่าที่แก้ได้ (เช่น α, ค่าสมมติใน Solver) | "
        "ตัวอักษรสีดำ = สูตร | พื้นเขียวอ่อน = ผลสรุป",
        "ทุกค่าเป็นสูตรที่อ้างอิงชีต Data ถ้าแก้หรือเพิ่มข้อมูลใน Data ผลทุกชีตจะคำนวณใหม่เอง "
        "(ถ้าเพิ่มแถว ให้สร้างไฟล์ใหม่จากเว็บ เพื่อขยายช่วงอ้างอิง)",
        "ระดับนัยสำคัญ α = 0.05 (แก้ได้ในแต่ละชีต) | p-value < α แปลว่าปฏิเสธ H₀",
        "ค่าสมมติทางธุรกิจ: กำไรต่อคนที่ตอบรับ 100 และต้นทุนต่อสาย 20 (ค่าเดียวกับโมเดล ML) "
        "งบจำนวนสายและโทรขั้นต่ำในชีต Solver เป็นค่าสมมติ",
        "ตรวจผลด้วย Data Analysis ToolPak ได้: Data → Data Analysis → Descriptive Statistics / "
        "Anova: Single Factor / t-Test: Two-Sample Assuming Unequal Variances",
        "ที่มาข้อมูล: UCI Machine Learning Repository – Bank Marketing (bank-additional.csv) "
        "ผ่านการบันทึก/แก้ไขในเว็บ Bank Call Tracker",
    ]
    _cell(ws, r, 1, "หมายเหตุ", font=F_BOLD, border=False)
    for i, n in enumerate(notes, start=r + 1):
        c = _cell(ws, i, 1, f"• {n}", border=False)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 30
    _widths(ws, [16, 70, 60])


# ============================================================
# ฟังก์ชันหลัก
# ============================================================
def build_analysis_workbook(customers: pd.DataFrame, group_name: str, generated_at: str) -> bytes:
    """customers = ตารางลูกค้าจากเว็บ (ใช้เฉพาะแถวที่ y เป็น yes/no) คืนไฟล์ .xlsx เป็น bytes"""
    df = customers[customers[C.TARGET_COL].isin(["yes", "no"])].copy()
    df = df[[C.ID_COL] + C.DATA_COLUMNS].reset_index(drop=True)
    for col in C.INT_COLUMNS:
        df[col] = pd.to_numeric(df[col]).astype(int)
    for col in C.MACRO_COLUMNS:
        df[col] = pd.to_numeric(df[col]).astype(float)
    for col in C.ALLOWED:
        df[col] = df[col].astype(str)
    if len(df) < 30:
        raise ValueError("มีลูกค้าที่โทรแล้วน้อยเกินไป (ต้องมีอย่างน้อย 30 คน)")

    wb = Workbook()
    _sheet_readme(wb, len(df), group_name, generated_at)
    refs = _sheet_data(wb, df)
    _sheet_descriptive(wb, refs)
    _sheet_conversion(wb, refs, df)
    _sheet_ttest(wb, refs)
    _sheet_anova(wb, refs, df)
    _sheet_chi(wb, refs, df)
    _sheet_ztest(wb, refs)
    _sheet_corr(wb, refs)
    _sheet_solver(wb, refs)
    wb.move_sheet("Data", offset=len(wb.sheetnames))        # ย้าย Data ไปท้ายสุด ให้ชีตผลอยู่ด้านหน้า
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = ws.title == "Data"
    wb.calculation.fullCalcOnLoad = True                       # ให้ Excel คำนวณสูตรทั้งหมดตอนเปิด
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
